# -*- coding: utf-8 -*-
"""Lap rap toan bo bao cao Word."""
import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import docxlib as D
import c0_front, c1_modau, c2_chuong1, c3_chuong2, c4_chuong3, c5_chuong4, c6_ketluan

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..",
                   "BAO_CAO_DE_TAI_SECURE_VAULT_RESEARCH.docx")


def toc_field(doc, levels="1-3"):
    """Chen truong muc luc tu dong (nhan F9 trong Word de cap nhat)."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run()._r
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    f1.set(qn("w:dirty"), "true")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "%s" \\h \\z \\u ' % levels
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Nhấn Ctrl+A rồi F9 trong Microsoft Word để cập nhật mục lục."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    r.append(f1); r.append(it); r.append(f2); r.append(t); r.append(f3)


def collect_indexes(all_blocks):
    figs, tbls = [], []
    for b in all_blocks:
        if b[0] in ("fig", "shot"):
            figs.append(b[2])
        elif b[0] == "tbl" and b[1]:
            tbls.append(b[1])
        elif b[0] == "ph":
            figs.append(b[1])
    return figs, tbls


def main():
    doc = D.new_document()

    body = (c1_modau.blocks() + c2_chuong1.blocks() + c3_chuong2.blocks()
            + c4_chuong3.blocks() + c5_chuong4.blocks() + c6_ketluan.blocks())
    figs, tbls = collect_indexes(body)

    # --- phan dau ---------------------------------------------------------
    D.render(doc, c0_front.blocks())

    # --- muc luc ----------------------------------------------------------
    doc.add_page_break()
    D.heading(doc, "MỤC LỤC", 1)
    toc_field(doc)

    # --- danh muc hinh ve -------------------------------------------------
    doc.add_page_break()
    D.heading(doc, "DANH MỤC HÌNH VẼ", 1)
    for f in figs:
        p = D.para(doc, f, indent=False, size=12, space_after=2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)

    # --- danh muc bang ----------------------------------------------------
    doc.add_page_break()
    D.heading(doc, "DANH MỤC BẢNG BIỂU", 1)
    for t in tbls:
        p = D.para(doc, t, indent=False, size=12, space_after=2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)

    # --- than bao cao -----------------------------------------------------
    D.render(doc, body)

    doc.save(OUT)
    print("Da luu:", OUT)
    print("So hinh:", len(figs), "| So bang:", len(tbls))
    words = 0
    for para in doc.paragraphs:
        words += len(para.text.split())
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                words += len(c.text.split())
    print("Uoc luong so tu:", words)
    print("So doan:", len(doc.paragraphs), "| So bang Word:", len(doc.tables))


if __name__ == "__main__":
    main()
